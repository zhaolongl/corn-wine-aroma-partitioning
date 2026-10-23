"""JAR merged paper, sixth version: supplementary-table workbook (Table S1-S18, numbered by first citation).
Unchanged tables are copied cell by cell (values, styles, widths, merges) from the fifth-version workbook and renumbered;
new tables are written from the v6 analysis outputs.  Run: python supp_tables.py <out.xlsx>"""
import copy
import os
import pickle
import sys

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as CL

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import data_merged as J                                                                # noqa: E402
import core as C                                                                 # noqa: E402

V5BOOK = str(C.CFG.RAW / "legacy_supplementary_tables_v5.xlsx")
MDRUN = str(C.CFG.RAW / "md")
MOD = pickle.load(open(os.path.join(C.OUT, "models.pkl"), "rb"))
EXT = pickle.load(open(os.path.join(C.OUT, "extra.pkl"), "rb"))
DK = pickle.load(open(os.path.join(C.OUT, "dock.pkl"), "rb"))
CUT = pickle.load(open(os.path.join(C.OUT, "cut.pkl"), "rb"))
HFILL = PatternFill("solid", fgColor="E9E4DA")
TITLES = {}


def copy_sheet(src, wb, name, new_title, new_note=None):
    ws = wb.create_sheet(name)
    for row in src.iter_rows():
        for c in row:
            n = ws.cell(row=c.row, column=c.column, value=c.value)
            if c.has_style:
                n.font = copy.copy(c.font); n.fill = copy.copy(c.fill); n.alignment = copy.copy(c.alignment)
                n.number_format = c.number_format; n.border = copy.copy(c.border)
    for k, d in src.column_dimensions.items():
        ws.column_dimensions[k].width = d.width
    for mr in src.merged_cells.ranges:
        ws.merge_cells(str(mr))
    ws["A1"] = new_title
    if new_note:
        ws["A2"] = new_note
    TITLES[name] = new_title
    return ws


def write_sheet(wb, name, title, note, df, nd=None, widths=None):
    ws = wb.create_sheet(name)
    ws["A1"] = title; ws["A1"].font = Font(bold=True, size=11)
    ws["A2"] = note; ws["A2"].alignment = Alignment(wrap_text=False)
    for j, col in enumerate(df.columns, 1):
        c = ws.cell(row=4, column=j, value=str(col)); c.font = Font(bold=True); c.fill = HFILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, row in enumerate(df.itertuples(index=False), 5):
        for j, v in enumerate(row, 1):
            if isinstance(v, (np.floating, float)):
                v = None if v != v else float(v)
            elif isinstance(v, (np.integer,)):
                v = int(v)
            elif isinstance(v, (np.bool_,)):
                v = bool(v)
            c = ws.cell(row=i, column=j, value=v)
            if isinstance(v, float):
                c.number_format = (nd or {}).get(df.columns[j - 1], "0.000")
    for j, col in enumerate(df.columns, 1):
        w = (widths or {}).get(col, max(9, min(28, len(str(col)) * 1.6 + 2)))
        ws.column_dimensions[CL(j)].width = w
    ws.freeze_panes = "A5"
    TITLES[name] = title
    return ws


def build(path):
    old = load_workbook(V5BOOK)
    wb = Workbook(); wb.remove(wb.active)
    info = wb.create_sheet("说明")
    # S1-S5 unchanged
    copy_sheet(old["Table S1"], wb, "Table S1", "Table S1. 化合物鉴定信息的核查与更正")
    copy_sheet(old["Table S2"], wb, "Table S2", "Table S2. 15 种共发酵酒的理化指标、酚类物质与抗氧化能力")
    copy_sheet(old["Table S3"], wb, "Table S3", "Table S3. 化学指标与感官属性的无重复双因素方差分解")
    copy_sheet(old["Table S4"], wb, "Table S4", "Table S4. 各化学指标随果汁比例的线性回归")
    D2 = pickle.load(open(os.path.join(C.CFG.RAW, "cofermentation_paper2_data.pkl"), "rb"))
    rw, pw, dfw = D2["corr_within"]
    chem = ["酒精度", "残糖", "总酸", "pH", "总酚", "总黄酮", "可溶性蛋白", "DPPH", "总还原能力"]
    sens = [c for c in rw.columns if c not in chem]
    rows5 = []
    for c in chem:
        for s_ in sens:
            rows5.append(dict(化学指标=c, 感官属性=s_.replace("感官总分", "QDA 强度总和"), r=float(rw.loc[c, s_]), P=float(pw.loc[c, s_])))
    t5 = pd.DataFrame(rows5)
    write_sheet(wb, "Table S5", "Table S5. 15 种共发酵酒化学指标与 QDA 属性的水果内 Pearson 相关",
                f"各指标先减去所属水果的均值后计算（n = 15，自由度 {dfw}）；这些相关主要反映各指标沿果汁比例梯度的共变，不代表因果关系。"
                "正文引用的澄清度、整体平衡性与总酚、总黄酮、总还原能力的相关见本表。", t5, nd={"r": "0.000", "P": "0.0000"})
    copy_sheet(old["Table S11"], wb, "Table S6", "Table S6. 共发酵酒 150 种挥发性化合物的含量、OAV、分子描述符与馏出类型预测")
    # S7 OAV robustness (new)
    w, bb = J.load(); kt = J.key_table(w, bb)
    oa = CUT["oav"]
    piv = oa.pivot_table(index=["CID"], columns=["体系", "浓度口径"], values="OAVmax", aggfunc="first")
    piv.columns = [f"OAVmax_{a}_{b}".replace("_—", "") for a, b in piv.columns]
    gr = oa.pivot_table(index=["CID"], columns=["体系", "浓度口径"], values="等级", aggfunc="first")
    gr.columns = [f"等级_{a}_{b}".replace("_—", "") for a, b in gr.columns]
    names = pd.concat([bb[["CID", "采用名称", "新类别", "阈值（μg/L）"]].rename(columns={"采用名称": "化合物", "新类别": "类别"}),
                       w[["CID", "化合物", "类别_统一", "阈值（μg/L）"]].rename(columns={"类别_统一": "类别"})]).drop_duplicates("CID")
    t7 = names.merge(piv.reset_index(), on="CID").merge(gr.reset_index(), on="CID")
    t7 = t7[(t7.filter(like="OAVmax").max(1) >= 0.5)].sort_values("OAVmax_馏出液_S0", ascending=False)
    t7["化合物"] = t7["化合物"].replace(C.REN)
    write_sheet(wb, "Table S7", "Table S7. 候选气味活性物质的阈值、最大 OAV 与阈值/浓度口径敏感性等级",
                "等级：稳定 OAVmax ≥ 2；阈值敏感 1 ≤ OAVmax < 2；临界 0.5 ≤ OAVmax < 1。S0 为表列浓度，S1 为稀释未回算条件情景；"
                "阈值来源与介质待逐一核对【待补】。仅列出任一体系 OAVmax ≥ 0.5 的化合物。", t7)
    st = CUT["stress"].pivot_table(index="阈值倍数", columns=["体系", "浓度口径"], values="候选数")
    st.columns = [f"{a}{'' if b == '—' else '（' + b + '）'}" for a, b in st.columns]
    ws = wb["Table S7"]; r0 = 5 + len(t7) + 2
    ws.cell(row=r0, column=1, value="阈值压力测试：各倍数下 OAVmax ≥ 1 的化合物数").font = Font(bold=True)
    for j, c in enumerate(["阈值倍数"] + list(st.columns), 1):
        ws.cell(row=r0 + 1, column=j, value=c).font = Font(bold=True)
    for i, (k, row) in enumerate(st.iterrows(), r0 + 2):
        ws.cell(row=i, column=1, value=float(k))
        for j, v in enumerate(row.values, 2):
            ws.cell(row=i, column=j, value=int(v))
    copy_sheet(old["Table S6"], wb, "Table S8", "Table S8. 二次蒸馏各段馏分与 5 个截取馏分的乙醇衡算")
    # S9 derived endpoints
    b = C.load()
    cols = ["编号", "采用名称", "CID", "新类别", "主分析集", "SMILES", "MW", "XLogP", "TPSA", "HBD", "HBA", "RotB", "芳香环", "O/C",
            "bp_exp", "洗脱类型", "n_det"] + C.FR + [f"p_{s}" for s in C.FRS] + ["R", "D", "H"] + \
           [f"p1_{s}" for s in C.FRS] + ["R_S1", "D_S1", "butina_main_0.6", "butina_all_0.6"]
    t9 = b[cols].rename(columns={"采用名称": "化合物", "新类别": "类别", "n_det": "检出馏分数", "bp_exp": "实验沸点（℃）",
                                 "butina_main_0.6": "Butina 簇（主集）", "butina_all_0.6": "Butina 簇（全集）"})
    write_sheet(wb, "Table S9", "Table S9. 115 种馏出液化合物的分子描述符、馏分浓度、质量份额与连续分配端点（S0 与 S1）",
                "p_j = c_j V_j / Σ c_j V_j（V = 50、300、50、100、200 mL）；R = p_中间 + p_次；D = Σ p_j t_j，t_j 为收集区间中点/700 mL；"
                "p1/R_S1/D_S1 为 S1 条件情景（前 4 馏分乘 7.88、7.14、5.24、2.99）。主分析集剔除烷烃、烯烃与芳香烃。", t9)
    # S10 performance
    p = pd.concat([MOD["perf"], EXT["perf_extra"]], ignore_index=True)
    lab = {"random": "随机 5 折", "group": "结构分组 0.6", "group0.5": "结构分组 0.5", "group0.7": "结构分组 0.7",
           "all115_random": "全部 115 种，随机", "all115_group": "全部 115 种，结构分组"}
    p["终点"] = p["方案"].str.split("|").str[0]; p["验证方案"] = p["方案"].str.split("|").str[1].map(lambda s: lab.get(s, s))
    p["模型"] = p["模型"].map({"baseline": "均值基线", "ridge": "Ridge", "xgb": "XGBoost"})
    t10 = p[["终点", "验证方案", "模型", "MAE_mean", "MAE_sd", "RMSE_mean", "RMSE_sd", "R2_mean", "R2_sd", "rho_mean", "rho_sd"]].rename(
        columns={"MAE_mean": "MAE 均值", "MAE_sd": "MAE 标准差", "RMSE_mean": "RMSE 均值", "RMSE_sd": "RMSE 标准差",
                 "R2_mean": "折外 R² 均值", "R2_sd": "折外 R² 标准差", "rho_mean": "Spearman ρ 均值", "rho_sd": "Spearman ρ 标准差"})
    write_sheet(wb, "Table S10", "Table S10. 连续端点预测模型在全部验证方案下的折外性能（10 次重复的均值与标准差）",
                f"响应置换（Ridge，随机 5 折，300 次）：R 的 P = {MOD['perm']['R']['p']:.3f}，D 的 P = {MOD['perm']['D']['p']:.3f}。"
                "标准化与超参数均在训练折内确定；预测截断至 0–1。R_S1/D_S1 为 S1 浓度情景；多馏分检出子集为至少 2 个馏分检出的化合物。", t10)
    # S11 per-compound OOF
    m = MOD["meta"]
    t11 = m[["编号", "采用名称", "新类别", "late", "R", "D"]].rename(columns={"采用名称": "化合物", "新类别": "类别", "late": "晚馏出型"}).copy()
    for ep in ["R", "D"]:
        for sch in ["random", "group"]:
            for mo in ["ridge", "xgb"]:
                oof = MOD[(ep, sch)][mo]["oof"]
                t11[f"{ep}_{lab[sch]}_{mo}_均值"] = oof.mean(0); t11[f"{ep}_{lab[sch]}_{mo}_SD"] = oof.std(0, ddof=1)
    t11["与训练折最大相似度_随机（中位）"] = np.median(MOD["nnsim"]["random"], 0)
    t11["与训练折最大相似度_分组（中位）"] = np.median(MOD["nnsim"]["group"], 0)
    write_sheet(wb, "Table S11", "Table S11. 主分析集 96 种化合物的逐化合物折外预测与结构相似度",
                "折外预测为 10 次重复的均值与标准差；相似度为测试化合物与训练折化合物的最大 Morgan 指纹 Tanimoto 相似度（10 次重复的中位数）。", t11)
    copy_sheet(old["Table S7"], wb, "Table S12", "Table S12. 早/晚馏出二分类 XGBoost 模型各重复的性能（辅助比较）",
               f"连续端点同一随机划分下的 logistic 分类 AUC = {MOD['late_auc'].mean():.3f} ± {MOD['late_auc'].std(ddof=1):.3f}；"
               "PLS-DA（AUC 0.763）、XGBoost（0.688 ± 0.038）与 logistic 回归（0.760 ± 0.014）见 Fig. S8。")
    # S13 wine-only predictions
    wp = EXT["wine_pred"].rename(columns={"类别_统一": "类别", "预测类型": "二分类多数票", "一致性": "三模型一致性", "h": "杠杆值",
                                           "D_pred": "Ridge 预测 D"})
    wp["化合物"] = wp["化合物"].replace(C.REN)
    write_sheet(wb, "Table S13", "Table S13. 共发酵酒独有且处于适用域内的 57 种化合物的馏出位置候选预测",
                f"Ridge 以主分析集 96 种化合物全部训练（α = {EXT['wine_alpha']:.3g}）；这些化合物无实测蒸馏标签，预测不构成外部验证，"
                "也不能推断其酒心保留。", wp.sort_values("OAVmax", ascending=False))
    # S14 ablation, class-out, correlations, coefficients
    abl = pd.concat([MOD["abl"], EXT["nonacid_bp"].assign(n=57)], ignore_index=True)
    abl["模型"] = abl["模型"].map({"ridge": "Ridge", "xgb": "XGBoost"})
    t14 = abl[["子集", "n", "终点", "特征组", "模型", "MAE", "MAE_sd", "R2", "基线MAE", "相对基线降低"]].rename(
        columns={"MAE_sd": "MAE 标准差", "R2": "折外 R²", "基线MAE": "基线 MAE", "相对基线降低": "MAE 相对基线降低（%）"})
    ws = write_sheet(wb, "Table S14", "Table S14. 特征组消融、类别留出、描述符相关与 Ridge 系数稳定性",
                     "相对基线降低 = (1 − MAE/基线 MAE) × 100%（10 次重复均值）；同一子集内所有特征组共用相同划分。", t14)
    r0 = 5 + len(t14) + 2
    blocks = []
    co = MOD["classout"].copy(); co["模型"] = co["模型"].map({"baseline": "均值基线", "ridge": "Ridge", "xgb": "XGBoost"})
    blocks.append(("类别留出压力测试（训练集不含该类别）", co))
    blocks.append(("描述符与 D、R 的 Spearman 相关（各子集内 Benjamini–Hochberg 校正）", MOD["cor"]))
    rows = []
    for (sname, ep), bt in MOD["boot"].items():
        for j, d in enumerate(C.DESC):
            rows.append(dict(子集=sname, 终点=ep, alpha=bt["alpha"], 描述符=d, 标准化系数=bt["coef"][j], 下限95=bt["lo"][j],
                             上限95=bt["hi"][j], 符号一致比例=bt["same_sign"][j]))
    blocks.append(("标准化 Ridge 系数（2000 次自助）", pd.DataFrame(rows)))
    blocks.append(("XGBoost（D，随机方案）TreeSHAP 重要性（50 个外层模型）", EXT["shap_D"]))
    tf = MOD["thr_fold"]; fx = MOD["thr_fixed"]
    blocks.append(("TPSA 阈值规则：固定 37 Å² 与折内搜索", pd.DataFrame([dict(规则="固定 37 Å²", 阈值=37.0, 灵敏度=fx["sens"], 特异度=fx["spec"]),
                                                               dict(规则="折内搜索（50 折均值）", 阈值=tf["thr"].median(),
                                                                    灵敏度=tf["sens"].mean(), 特异度=tf["spec"].mean())])))
    for ttl, df in blocks:
        ws.cell(row=r0, column=1, value=ttl).font = Font(bold=True)
        for j, c in enumerate(df.columns, 1):
            ws.cell(row=r0 + 1, column=j, value=str(c)).font = Font(bold=True)
        for i, row in enumerate(df.itertuples(index=False), r0 + 2):
            for j, v in enumerate(row, 1):
                if isinstance(v, (np.floating, float)):
                    v = None if v != v else float(v)
                elif isinstance(v, np.integer):
                    v = int(v)
                ws.cell(row=i, column=j, value=v)
        r0 += len(df) + 4
    # S15 docking (old full table + new structural statistics)
    copy_sheet(old["Table S12"], wb, "Table S15", "Table S15. 扩展分子对接的全部评分（31 个配体 × 6 个受体 × 3 个随机种子）与结构关联统计",
               "评分为 Vina 最优构象评分（kcal/mol）；结构关联统计（种子均值、重原子数校正评分、结构项、分配关联）见本表右侧。")
    ws = wb["Table S15"]; c0 = 10
    lig = DK["lig"]
    keep = ["ligand", "code", "CID", "HA", "TPSA", "arom", "R", "D", "late"] + [f"v_{r}" for r in
                                                                                  ["OR1A1", "OR1A2", "OR1D2", "OR1G1", "OR2W1", "OR51E2"]] + \
           [f"c_{r}" for r in ["OR1A1", "OR1A2", "OR1D2", "OR1G1", "OR2W1", "OR51E2"]] + ["c_mean"]
    for (ttl, df) in [("配体层面：种子均值评分（v_）与重原子数校正评分（c_）", lig[keep]), ("结构项（校正评分 ~ z(TPSA) + z(芳香环)）", DK["struct"]),
                      ("分配端点与校正评分的关联（n = 30）", DK["assoc"]), ("早/晚馏出配体的校正评分比较", DK["group"])]:
        ws.cell(row=4 if ttl.startswith("配体") else ws.max_row + 3, column=c0, value=ttl).font = Font(bold=True)
        r = ws.max_row + 1 if not ttl.startswith("配体") else 5
        for j, c in enumerate(df.columns):
            ws.cell(row=r, column=c0 + j, value=str(c)).font = Font(bold=True)
        for i, row in enumerate(df.itertuples(index=False), r + 1):
            for j, v in enumerate(row):
                if isinstance(v, (np.floating, float)):
                    v = None if v != v else float(v)
                elif isinstance(v, np.integer):
                    v = int(v)
                ws.cell(row=i, column=c0 + j, value=v)
    # S16 MD (current trajectories)
    sm = pd.read_csv(os.path.join(MDRUN, "md_summary.csv")); ct = pd.read_csv(os.path.join(MDRUN, "md_contacts.csv"))
    ws = write_sheet(wb, "Table S16", "Table S16. OR1A2 复合物分子动力学模拟的轨迹统计与残基接触占比",
                     f"各复合物已完成 {sm['ns'].min():.0f} ns（计划 100 ns，模拟进行中【完成后更新】）；RMSD、RMSF 与相互作用能为生产模拟后半段的平均值；"
                     "残基编号为模拟编号（加 8 为 UniProt 编号）。", sm)
    r0 = 5 + len(sm) + 2
    ws.cell(row=r0, column=1, value="残基接触帧占比").font = Font(bold=True)
    for j, c in enumerate(ct.columns, 1):
        ws.cell(row=r0 + 1, column=j, value=c).font = Font(bold=True)
    for i, row in enumerate(ct.itertuples(index=False), r0 + 2):
        for j, v in enumerate(row, 1):
            ws.cell(row=i, column=j, value=(float(v) if isinstance(v, (np.floating, float)) else int(v) if isinstance(v, np.integer) else v))
    # S17 cut scenarios
    cb = CUT["combos"].copy()
    ws = write_sheet(wb, "Table S17", "Table S17. 截取情景：整馏分组合、掐头体积与头馏分内部分布情景下的乙醇与候选物保留",
                     "整馏分组合可由实测份额直接计算；头馏分内部截取依赖有标签的分布情景（前置富集 ∝ e^(−3x)、均匀、后置富集 ∝ e^(3x)），"
                     "情景不代表概率。判断以酒心占比 50% 为界。", cb)
    r0 = 5 + len(cb) + 2
    od = CUT["odor"].copy()
    for ttl, df in [("25 种馏出液候选物在三个掐头水平下的情景范围", od), ("乙醇回收率与掐头体积（头内均匀）", CUT["ethanol"])]:
        ws.cell(row=r0, column=1, value=ttl).font = Font(bold=True)
        for j, c in enumerate(df.columns, 1):
            ws.cell(row=r0 + 1, column=j, value=str(c)).font = Font(bold=True)
        for i, row in enumerate(df.itertuples(index=False), r0 + 2):
            for j, v in enumerate(row, 1):
                if isinstance(v, (np.floating, float)):
                    v = None if v != v else float(v)
                elif isinstance(v, np.integer):
                    v = int(v)
                ws.cell(row=i, column=j, value=v)
        r0 += len(df) + 4
    copy_sheet(old["Table S9"], wb, "Table S18", "Table S18. 截取与预热工艺的单因素试验与 L9(3⁴) 正交试验",
               "单因素试验：均值 ± 标准差（n = 3），不同小写字母表示同一因素内差异显著（P < 0.05）。正交试验与方差分析见本表下方。")
    ws = wb["Table S18"]; r0 = ws.max_row + 3
    L9 = [["试验号", "A 掐头去尾量（%）", "B 预热时间（min）", "C 预热温度（℃）", "感官评分（分）"],
          [1, 44, 7, 68, 73.5], [2, 44, 10, 72, 77.2], [3, 44, 13, 76, 68.9], [4, 45, 7, 72, 82.8], [5, 45, 10, 76, 79.4],
          [6, 45, 13, 68, 75.0], [7, 46, 7, 76, 70.4], [8, 46, 10, 68, 69.4], [9, 46, 13, 72, 73.5],
          ["k1", 73.20, 75.57, 72.63, None], ["k2", 79.07, 75.33, 77.83, None], ["k3", 71.10, 72.47, 72.90, None],
          ["R", 7.97, 3.10, 5.20, None], ["平方和", 102.30, 17.88, 51.45, "误差 4.98（df = 2）"], ["F", 20.56, 3.59, 10.34, None],
          ["P", 0.046, 0.218, 0.088, "最优组合 A2B1C2；各处理测定 1 次，检验效能有限；验证试验 82.8 ± 1 分【待核实】"]]
    ws.cell(row=r0 - 1, column=1, value="L9(3⁴) 正交试验、极差分析与方差分析").font = Font(bold=True)
    for i, row in enumerate(L9):
        for j, v in enumerate(row, 1):
            c = ws.cell(row=r0 + i, column=j, value=v)
            if i == 0:
                c.font = Font(bold=True)
    # info sheet
    info["A1"] = "论文合并稿（Journal of Advanced Research）第六版补充材料：Table S1–S18（按正文首次引用顺序编号）"
    info["A1"].font = Font(bold=True, size=12)
    info["A2"] = ("计算分析由 Python 3.11（scikit-learn 1.3.0、xgboost 3.0.1、SciPy、statsmodels）与 RDKit 2026.09 完成；"
                  "对接与动力学由 AutoDock Vina 1.2.7、PLIP 与 OpenMM 8.6 完成。Table S1–S4、S6、S8、S12、S15（评分部分）与 S18 为共发酵与蒸馏试验的已核对数据。")
    for i, (k, v) in enumerate(TITLES.items(), 4):
        info.cell(row=i, column=1, value=k); info.cell(row=i, column=2, value=v)
    info.column_dimensions["A"].width = 12; info.column_dimensions["B"].width = 110
    wb.save(path)
    print("saved", path, len(wb.sheetnames), "sheets")


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else os.path.join(C.CFG.TABLES, "Supplementary_Tables_S1-S18.xlsx"))
